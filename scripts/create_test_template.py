"""Synthetic automated-test fixture. Never a church original or production template."""
from pathlib import Path
from docx import Document
root=Path(__file__).resolve().parents[1]
(root/'artifacts').mkdir(exist_ok=True)
doc=Document();doc.add_paragraph('SYNTHETIC TEST ONLY · 순모임 보고서')
t=doc.add_table(rows=2,cols=8);t.style='Table Grid'
for c,s in zip(t.rows[0].cells,['이름','주일예배','대학부예배','순모임','큐티','설교나눔','삶나눔','기도제목']):c.text=s
doc.add_paragraph('TO. 목사님');doc.save(root/'artifacts/synthetic-test.docx')
