import asyncio, io, json, time
from pathlib import Path
from unittest.mock import patch
import pytest
from docx import Document
from pydantic import ValidationError
from backend.app.schemas import Report
from backend.app.services.ai_analysis import analyze,validate_people
from backend.app.services.docx_generator import generate
from backend.app.services.transcript import parse_transcript
from backend.app.main import folder
H={'X-Moimlog':'1'}
TEXT='Speaker 1 [00:00:01]: 학교 시험이 끝났어요.\nSpeaker 2 [00:00:05]: 취업 준비 중입니다.\nSpeaker 3 [00:00:08]: 오늘의 말씀입니다.'

def upload(c,m,text=TEXT):
    return c.post(f'/api/meetings/{m["id"]}/upload',files={'file':('meeting.txt',text.encode(),'text/plain')},headers=H)

def map_speakers(c,m,mapping=None):
    return c.put(f'/api/meetings/{m["id"]}/speakers',json={'mapping':mapping or {'Speaker 1':'p1','Speaker 2':'p2','Speaker 3':'OTHER'}},headers=H)

def local(c,m):
    upload(c,m);map_speakers(c,m)
    r=c.post(f'/api/meetings/{m["id"]}/analyze',json={'mode':'local'},headers=H);assert r.status_code==200
    for _ in range(100):
        d=c.get('/api/meetings/'+m['id']).json()
        if d['job']['state']!='processing':return d
        time.sleep(.01)
    pytest.fail('job did not finish')

def template_bytes():
    doc=Document();doc.add_paragraph('순모임 보고서')
    meta=doc.add_table(rows=2,cols=2);meta.cell(0,0).text='모임 날짜';meta.cell(1,0).text='순장'
    t=doc.add_table(rows=2,cols=8);t.style='Table Grid'
    for c,s in zip(t.rows[0].cells,['이름','주일예배','대학부예배','순모임','큐티','설교나눔','삶나눔','기도제목']):c.text=s
    t.rows[1].height=600000
    for c in t.rows[1].cells:c.paragraphs[0].add_run('').font.size=120000
    doc.add_paragraph('TO. 목사님')
    out=io.BytesIO();doc.save(out);return out.getvalue()

def test_auth_required_and_csrf(client,info):
    assert client.get('/api/meetings').status_code==401
    assert client.post('/api/meetings',json=info).status_code==403

def test_txt_upload(account,meeting):
    r=upload(account,meeting);assert r.status_code==200
    d=r.json();assert d['text']==TEXT;assert len(d['segments'])==3;assert d['step']==3

def test_txt_encoding_and_empty(account,meeting):
    r=account.post(f'/api/meetings/{meeting["id"]}/upload',files={'file':('x.txt',b'\xff\xfe','text/plain')},headers=H)
    assert r.status_code==400
    assert upload(account,meeting,'  ').status_code==400

def test_upload_limits_and_extension(account,meeting):
    r=account.post(f'/api/meetings/{meeting["id"]}/upload',files={'file':('x.txt',b'a'*(2*1024*1024+1),'text/plain')},headers=H)
    assert r.status_code==413
    r=account.post(f'/api/meetings/{meeting["id"]}/upload',files={'file':('../../x.exe',b'bad','application/octet-stream')},headers=H)
    assert r.status_code==400

def test_participants_reorder_remove_duplicate_ids(account,meeting,info):
    info['participants'].reverse();r=account.patch('/api/meetings/'+meeting['id'],json=info,headers=H)
    assert r.json()['participants'][0]['id']=='p2'
    info['participants']=info['participants'][:1];assert account.patch('/api/meetings/'+meeting['id'],json=info,headers=H).status_code==200
    info['participants']*=2;assert account.patch('/api/meetings/'+meeting['id'],json=info,headers=H).status_code==422

def test_mapping_change_and_invalid_ids(account,meeting):
    upload(account,meeting);assert map_speakers(account,meeting).status_code==200
    assert map_speakers(account,meeting,{'Speaker 1':'p2'}).json()['mapping']['Speaker 1']=='p2'
    assert map_speakers(account,meeting,{'Speaker 999':'p1'}).status_code==422
    assert map_speakers(account,meeting,{'Speaker 1':'nonexistent'}).status_code==422

def test_local_no_invented_sermon_or_prayer_and_same_names(account,meeting):
    d=local(account,meeting);assert d['job']['state']=='done'
    rows=d['report']['participants'];assert rows[0]['name']==rows[1]['name']
    assert '학교 시험' in rows[0]['lifeSharing'];assert '취업 준비' not in rows[0]['lifeSharing']
    assert rows[0]['sermonSharing']=='';assert rows[0]['prayerRequests']=='';assert rows[0]['prayerSource']=='none'
    assert all(p['sunday']=='unknown' and p['meeting']=='unknown' for p in rows)
    assert d['mode']=='local'

def test_plain_txt_remains_unassigned(account,meeting):
    upload(account,meeting,'이름을 모르는 이야기입니다.');r=account.post('/api/meetings/'+meeting['id']+'/analyze',json={'mode':'local'},headers=H)
    for _ in range(100):
        d=account.get('/api/meetings/'+meeting['id']).json()
        if d['job']['state']=='done':break
        time.sleep(.01)
    assert d['report']['unassignedStatements'];assert all(not p['lifeSharing'] for p in d['report']['participants'])

def test_report_edit_save_complete_and_refresh(account,meeting):
    d=local(account,meeting);r=d['report'];r['participants'][0]['sermonSharing']='수정한 설교나눔';r['participants'][0]['meeting']='yes'
    out=account.put('/api/meetings/'+meeting['id']+'/report',json={'report':r,'complete':True},headers=H)
    assert out.status_code==200;assert out.json()['status']=='complete'
    stored=account.get('/api/meetings/'+meeting['id']+'/report').json()
    assert stored['participants'][0]['sermonSharing']=='수정한 설교나눔';assert stored['participants'][0]['meeting']=='yes'

def test_json_validation_and_incorrect_participant_ids(account,meeting):
    with pytest.raises(ValidationError):Report.model_validate_json('{broken')
    with pytest.raises(ValidationError):Report.model_validate({'meetingDate':'2026-01-01'})
    d=local(account,meeting);r=Report.model_validate(d['report']);r.participants[0].participantId='fake'
    with pytest.raises(ValueError):validate_people(r,d['participants'])
    assert account.put('/api/meetings/'+meeting['id']+'/report',json={'report':r.model_dump()},headers=H).status_code==422

def test_ai_no_key_and_consent(account,meeting):
    upload(account,meeting)
    assert account.post('/api/meetings/'+meeting['id']+'/analyze',json={'mode':'ai','consent':True},headers=H).status_code==400

def test_ai_failure_retries_without_exposing_secrets(monkeypatch):
    monkeypatch.setenv('OPENAI_API_KEY','secret-test')
    class Failing:
        calls=0
        def __init__(self,**kw):self.responses=self
        async def parse(self,**kw):Failing.calls+=1;raise RuntimeError('secret-test transcript must not leak')
    with patch('backend.app.services.ai_analysis.AsyncOpenAI',Failing):
        with pytest.raises(ValueError,match='원문은 보존') as e:
            asyncio.run(analyze({'participants':[{'id':'p1','name':'김'}],'segments':[{'speaker':'S','text':'원문'}],'mapping':{}},'ai'))
    assert Failing.calls==2;assert 'secret-test' not in str(e.value)

def test_ai_structured_response_missing_and_unknown_ids_retry(monkeypatch):
    monkeypatch.setenv('OPENAI_API_KEY','test')
    class Invalid:
        calls=0
        def __init__(self,**kw):self.responses=self
        async def parse(self,**kw):
            Invalid.calls+=1
            return type('Response',(),{'output_parsed':None})()
    with patch('backend.app.services.ai_analysis.AsyncOpenAI',Invalid):
        with pytest.raises(ValueError):asyncio.run(analyze({'participants':[],'segments':[{'text':'a','speaker':None}]},'ai'))
    assert Invalid.calls==2

def test_missing_template(account,meeting,tmp_path):
    local(account,meeting)
    with patch('backend.app.main.template_path',return_value=tmp_path/'missing.docx'):
        r=account.post('/api/meetings/'+meeting['id']+'/generate-docx',headers=H)
    assert r.status_code==422;assert '원본' in r.json()['detail']

def test_docx_generate_and_binary_download(account,meeting,tmp_path):
    d=local(account,meeting)
    original=template_bytes();r=account.post('/api/template',files={'file':('26_순보고서.docx',original)},headers=H);assert r.status_code==200
    r=account.post('/api/meetings/'+meeting['id']+'/generate-docx',headers=H);assert r.status_code==200
    downloaded=account.get(r.json()['url']);assert downloaded.content[:2]==b'PK'
    assert downloaded.headers['content-type']=='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    assert '2026-10-18_moim_report.docx' in downloaded.headers['content-disposition']
    doc=Document(io.BytesIO(downloaded.content));assert len(doc.tables[1].rows)==3
    assert doc.tables[1].rows[1].cells[0].text=='김정환';assert '시험' in doc.tables[1].rows[1].cells[6].text
    assert doc.tables[0].cell(0,1).text=='2026-10-18';assert doc.paragraphs[-1].text=='TO. 목사님'
    assert doc.tables[1].rows[1].height==Document(io.BytesIO(original)).tables[1].rows[1].height
    assert account.get(r.json()['url']).status_code==200

def test_template_bad_zip_and_unknown_layout(account,meeting):
    assert account.post('/api/template',files={'file':('x.docx',b'fake')},headers=H).status_code==400
    local(account,meeting);doc=Document();doc.add_table(rows=2,cols=2);out=io.BytesIO();doc.save(out)
    assert account.post('/api/template',files={'file':('x.docx',out.getvalue())},headers=H).status_code==200
    assert account.get('/api/meetings/'+meeting['id']+'/download').status_code==422

def test_original_repeated_block_expansion(tmp_path):
    doc=Document();t=doc.add_table(rows=5,cols=4);t.style='Table Grid'
    t.rows[0].cells[0].text='순장';t.rows[0].cells[2].text='모임 날짜'
    for row,label in zip(t.rows[1:],['이름','설교나눔','삶나눔','기도제목']):row.cells[0].text=label;row.cells[1].merge(row.cells[3])
    doc.add_paragraph('TO. 목사님');path=tmp_path/'original.docx';doc.save(path)
    people=[{'name':str(i),'sermonSharing':'말씀','lifeSharing':'삶','prayerRequests':''} for i in range(6)]
    data={'meetingDate':'2026-10-18','leader':'김','report':{'participants':people,'meetingSummary':''}}
    result=Document(io.BytesIO(generate(path,data)))
    assert len(result.tables[0].rows)==25;assert result.tables[0].rows[21].cells[1].text=='5';assert result.paragraphs[-1].text=='TO. 목사님'

def test_cross_account_access_all_resources(account,meeting):
    local(account,meeting);account.post('/api/auth/logout',headers=H)
    account.post('/api/auth/register',json={'email':str(time.time())+'@other.test','password':'another-password'},headers=H)
    base='/api/meetings/'+meeting['id']
    for path in ['', '/report','/transcription','/download']:
        assert account.get(base+path).status_code==404
    assert account.delete(base,headers=H).status_code==404
    assert account.post(base+'/generate-docx',headers=H).status_code==404
    assert account.post(base+'/upload',files={'file':('x.txt',b'text')},headers=H).status_code==404
    assert account.get('/api/meetings').json()==[]

def test_delete_report_and_account(account,meeting):
    assert account.delete('/api/meetings/'+meeting['id'],headers=H).status_code==200
    assert account.get('/api/meetings/'+meeting['id']).status_code==404
    assert account.delete('/api/me',headers=H).status_code==200
    assert account.get('/api/me').status_code==401

def test_transcript_timestamp_formats():
    s=parse_transcript('Speaker 1 [00:01:20]\n학교 이야기\nSpeaker 2 [00:01:27]: 기도해주세요')
    assert s[0]['speaker']=='Speaker 1';assert s[0]['timestamp']=='00:01:20';assert '학교' in s[0]['text']

def test_supplied_template_preserves_layout_and_expands():
    import unicodedata
    template=next((p for p in Path('backend/templates').glob('*.docx') if unicodedata.normalize('NFC',p.name)=='26_순보고서.docx'),None)
    if not template:pytest.skip('교회 원본 양식이 제공되지 않았습니다.')
    original=Document(template)
    people=[{'name':f'순원{i}', 'sunday':'yes','youth':'no','meeting':'unknown','qt':'yes','sermonSharing':f'말씀 {i}','lifeSharing':f'삶 {i}','prayerRequests':f'기도 {i}'} for i in range(9)]
    data={'meetingDate':'2026-10-18','leader':'김순장','report':{'participants':people,'meetingSummary':''}}
    generated=Document(io.BytesIO(generate(template,data)))
    assert original.sections[0]._sectPr.xml==generated.sections[0]._sectPr.xml
    assert len(generated.tables)==2 and len(generated.tables[1].rows)==8
    for before,after in zip(original.tables,generated.tables):
        assert before._tbl.tblPr.xml==after._tbl.tblPr.xml
        assert before.rows[0]._tr.xml==after.rows[0]._tr.xml
        assert before.rows[1]._tr.xml==after.rows[1]._tr.xml
    assert '2026년  10월  18일' in generated.paragraphs[1].text and '김순장' in generated.paragraphs[1].text
    row=generated.tables[1].rows[6]
    assert [c.text for c in row.cells[:5]]==['순원8','O','X','미확인','O']
    assert all(s in row.cells[5].text for s in ['말씀 8','삶 8','기도 8'])
    assert generated.tables[1].rows[-1].cells[0].text==original.tables[1].rows[-1].cells[0].text

def test_supplied_template_download_endpoint(account,meeting):
    from backend.app.main import template_path
    if not template_path({'id':'unused-test-user'}).exists():pytest.skip('교회 원본 양식이 제공되지 않았습니다.')
    local(account,meeting)
    response=account.get('/api/meetings/'+meeting['id']+'/download')
    assert response.status_code==200
    doc=Document(io.BytesIO(response.content))
    assert doc.tables[0].rows[2].cells[0].text=='김정환'
    assert '학교 시험' in doc.tables[0].rows[2].cells[5].text
