"""Regression coverage against the supplied, unmodified church form."""
from io import BytesIO
from pathlib import Path
import unicodedata

import pytest
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE

from backend.app.services.docx_generator import generate


TEMPLATE = next(p for p in Path('backend/templates').glob('*.docx')
                if unicodedata.normalize('NFC', p.name) == '26_순보고서.docx')


@pytest.mark.parametrize('count', [6, 8, 10])
@pytest.mark.parametrize('long_text', [False, True])
def test_original_report(count, long_text):
    original = Document(TEMPLATE)
    people = [dict(name=f'순원{i}', sunday='yes', youth='no', meeting='unknown', qt='yes',
                   sermonSharing=f'설교{i} ' + ('말씀과 적용을 나누었습니다. ' * 100 if long_text else '말씀'),
                   lifeSharing=f'삶{i} ' + ('일상과 어려움을 나누었습니다. ' * 100 if long_text else '일상'),
                   prayerRequests=f'기도{i} ' + ('가족을 위해 기도합니다. ' * 100 if long_text else '가족'))
              for i in range(count)]
    data = {'meetingDate': '2026-10-18', 'leader': '김순장',
            'report': {'participants': people, 'meetingSummary': ''}}
    result = Document(BytesIO(generate(TEMPLATE, data)))
    assert result.sections[0]._sectPr.xml == original.sections[0]._sectPr.xml
    assert len(result.tables) == 2
    assert [len(t.rows) for t in result.tables] == [6, 7 + max(0, count - 8)]
    for before, after in zip(original.tables, result.tables):
        assert before._tbl.tblPr.xml == after._tbl.tblPr.xml
        assert before._tbl.tblGrid.xml == after._tbl.tblGrid.xml
        for i in (0, 1):
            assert before.rows[i]._tr.xml == after.rows[i]._tr.xml
    assert result.tables[1].rows[-1]._tr.xml == original.tables[1].rows[-1]._tr.xml
    assert '2026년  10월  18일' in result.paragraphs[1].text
    assert '순 장 :  김순장' in result.paragraphs[1].text
    rows = list(result.tables[0].rows)[2:] + list(result.tables[1].rows)[2:-1]
    for person, row in zip(people, rows):
        assert [c.text for c in row.cells[:5]] == [person['name'], 'O', 'X', '미확인', 'O']
        paragraphs = row.cells[5].paragraphs
        for label, key in [('1.설교나눔:', 'sermonSharing'), ('2.삶나눔:', 'lifeSharing'), ('3.기도제목:', 'prayerRequests')]:
            index = next(i for i, p in enumerate(paragraphs) if ''.join(p.text.split()) == label)
            assert paragraphs[index + 1].text == person[key]
        assert row.height_rule == WD_ROW_HEIGHT_RULE.AT_LEAST
        assert not row._tr.xpath('./w:trPr/w:cantSplit')
