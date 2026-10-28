"""Fill an original DOCX, preserving XML formatting; never manufacture a template.
Supported layouts: tabular participant columns or repeating label-based blocks.
Unknown layouts fail closed with an actionable message.
"""
from copy import deepcopy
from io import BytesIO
import re
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.table import _Row

LABELS={'이름':'name','주일예배':'sunday','대학부예배':'youth','순모임':'meeting','큐티':'qt','설교나눔':'sermonSharing','삶나눔':'lifeSharing','기도제목':'prayerRequests'}
def norm(s):return re.sub(r'\s','',s)
def fill(cell,value):
    # Keep cell properties, paragraph properties and first run font formatting.
    p=cell.paragraphs[0]
    run=p.runs[0] if p.runs else p.add_run()
    run.text=str(value)
    for other in list(p.runs)[1:]:other._element.getparent().remove(other._element)
    for other in list(cell.paragraphs)[1:]:other._element.getparent().remove(other._element)

def value(person,key):
    v=person.get(key,'')
    if key in ['sunday','youth','meeting','qt']:return {'yes':'O','no':'X','unknown':'미확인'}.get(v,'미확인')
    return v

def original_layout(document, data):
    """Fill the supplied 2026 form's two six-column tables in place."""
    tables=document.tables
    def matches(table):
        return (len(table.rows)>=3 and len(table.columns)==6
                and '항목' in norm(table.cell(0,0).text)
                and '설교나눔' in norm(table.cell(0,5).text)
                and '주일예배' in norm(table.cell(1,1).text))
    if len(tables)!=2 or not all(matches(t) for t in tables):return False
    slots=[]
    for table in tables:
        for row in list(table.rows)[2:]:
            if '1.설교나눔:' in norm(row.cells[5].text):slots.append((table,row))
    if not slots:return False
    people=data['report']['participants']
    last_table,last_row=slots[-1]
    while len(slots)<len(people):
        element=deepcopy(last_row._tr)
        last_row._tr.addnext(element)
        last_row=_Row(element,last_table)
        slots.append((last_table,last_row))
    for person,(_,row) in zip(people,slots):
        for col,key in enumerate(['name','sunday','youth','meeting','qt']):
            fill(row.cells[col],value(person,key))
        paragraphs=row.cells[5].paragraphs
        for label,key in [('1.설교나눔:','sermonSharing'),('2.삶나눔:','lifeSharing'),('3.기도제목:','prayerRequests')]:
            index=next(i for i,p in enumerate(paragraphs) if norm(p.text).startswith(label))
            # Write into the existing blank paragraph, retaining label and paragraph styles.
            p=paragraphs[index+1]
            run=p.runs[0] if p.runs else p.add_run()
            run.text=value(person,key)
        # Fixed template heights clip long text; retain minimum height and allow expansion.
        row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
    for p in document.paragraphs:
        if '순 장' in p.text and '월' in p.text and '일' in p.text:
            year,month,day=data['meetingDate'].split('-')
            if len(p.runs)>=2:
                p.runs[0].text=f' {year}년  {int(month)}월  {int(day)}일'
                p.runs[1].text=f'\t순 장 :  {data["leader"]}'
            break
    return True

def generate(template,data):
    if not template.exists():raise ValueError('원본 26_순보고서.docx를 설정에서 등록해 주세요.')
    document=Document(template)
    if original_layout(document,data):
        out=BytesIO();document.save(out);return out.getvalue()
    report=data['report']; people=report['participants']; written=False
    # Replace explicit metadata placeholders without replacing entire paragraphs.
    replacements={'{{meetingDate}}':data['meetingDate'],'{{leader}}':data['leader'],'{{title}}':data.get('title',''),'{{meetingSummary}}':report['meetingSummary']}
    def replace(p):
        for token,v in replacements.items():
            if token in p.text:
                # tokens spanning runs: preserve the first run style
                text=p.text.replace(token,v)
                if p.runs:
                    p.runs[0].text=text
                    for r in p.runs[1:]:r.text=''
                else:p.add_run(text)
    for p in document.paragraphs:replace(p)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:replace(p)
        # Dates / leader label followed by value cell.
        for row in table.rows:
            seen=set()
            cells=[c for c in row.cells if not (id(c._tc) in seen or seen.add(id(c._tc)))]
            for i,c in enumerate(cells[:-1]):
                label=norm(c.text)
                if label in ['모임날짜','날짜','모임일시']:fill(cells[i+1],data['meetingDate'])
                if label=='순장':fill(cells[i+1],data['leader'])
        # One-row-per-person table with explicit column names.
        for index,row in enumerate(table.rows):
            headers={i:LABELS[norm(c.text)] for i,c in enumerate(row.cells) if norm(c.text) in LABELS}
            if set(headers.values()) >= {'name','sermonSharing','lifeSharing','prayerRequests'}:
                if index+1>=len(table.rows):raise ValueError('양식의 참석자 예시 행이 필요합니다.')
                prototype=deepcopy(table.rows[index+1]._tr)
                # Only consume blank or existing participant rows, preserving footer.
                existing=[]
                for r in list(table.rows)[index+1:]:
                    if any('TO.' in c.text or '목사님' in c.text for c in r.cells):break
                    existing.append(r)
                anchor=row._tr
                for person in people:
                    element=deepcopy(prototype);anchor.addnext(element);anchor=element
                    from docx.table import _Row
                    out=_Row(element,table)
                    for col,key in headers.items():fill(out.cells[col],value(person,key))
                for r in existing:r._tr.getparent().remove(r._tr)
                written=True;break
        if written:break
        # Repeated original block: name/attendance row followed by three labelled rows.
        rows=list(table.rows);starts=[i for i,r in enumerate(rows) if any(norm(c.text)=='이름' for c in r.cells)]
        for start in starts:
            end=None; found=set()
            for i in range(start,min(start+12,len(rows))):
                found|={norm(c.text) for c in rows[i].cells}
                if {'설교나눔','삶나눔','기도제목'}<=found:end=i+1;break
            if end is None:continue
            block=[deepcopy(r._tr) for r in rows[start:end]]
            anchor=rows[start]._tr
            for person in people:
                from docx.table import _Row
                for raw in block:
                    el=deepcopy(raw);anchor.addprevious(el)
                    r=_Row(el,table); unique=[]
                    for c in r.cells:
                        if not any(c._tc is x._tc for x in unique):unique.append(c)
                    for j,c in enumerate(unique[:-1]):
                        key=LABELS.get(norm(c.text))
                        if key and norm(unique[j+1].text) not in LABELS:fill(unique[j+1],value(person,key))
            # Remove all original blocks (leave unrelated footer rows untouched).
            remove=set()
            for s in starts:
                for r in rows[s:s+(end-start)]:remove.add(r._tr)
            for el in remove:
                if el.getparent() is not None:el.getparent().remove(el)
            written=True;break
        if written:break
    if not written:raise ValueError('양식 구조를 인식하지 못했습니다. 이름·설교나눔·삶나눔·기도제목 열 또는 반복 블록이 있는 원본 양식이 필요합니다. 관리자에게 양식 호환 작업을 요청해 주세요.')
    out=BytesIO(); document.save(out);return out.getvalue()
